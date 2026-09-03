"""Tests for token accounting and the web app's password gate.

No model calls and no network - the ledger is fed synthetic usage objects, and
the app is exercised through Flask's test client.

Run:  python test_usage_auth.py
"""

import base64
import json
import os
import sys
import tempfile

os.environ.setdefault("ANTHROPIC_API_KEY", "test-key-not-used")

import seo_writer as sw


# --- pricing ---------------------------------------------------------------

def test_known_model_is_priced_from_the_table():
    # claude-sonnet-5 is $2 in / $10 out per million.
    cost = sw._price("claude-sonnet-5", 1_000_000, 1_000_000)
    assert abs(cost - 12.0) < 1e-9, cost


def test_opus_costs_more_than_sonnet_for_the_same_work():
    assert sw._price("claude-opus-5", 100_000, 50_000) > \
           sw._price("claude-sonnet-5", 100_000, 50_000)


def test_unknown_model_has_no_price_rather_than_a_wrong_one():
    assert sw._price("some-model-we-never-heard-of", 1000, 1000) is None


def test_cache_reads_are_cheaper_than_fresh_input():
    fresh = sw._price("claude-sonnet-5", 1_000_000, 0)
    cached = sw._price("claude-sonnet-5", 0, 0, cache_read=1_000_000)
    assert cached < fresh, (cached, fresh)


# --- the ledger ------------------------------------------------------------

def test_usage_is_tallied_by_model_and_step():
    sw.reset_usage()
    sw.log("STEP 5", "writing")
    sw.record_usage("anthropic", "claude-sonnet-5", 1000, 500)
    sw.log("STEP 6.5", "auditing")
    sw.record_usage("anthropic", "claude-opus-5", 2000, 300)
    sw.record_usage("anthropic", "claude-opus-5", 1000, 200)

    s = sw.usage_summary()
    assert s["total"]["calls"] == 3
    assert s["total"]["input_tokens"] == 4000
    assert s["total"]["output_tokens"] == 1000
    assert s["total"]["total_tokens"] == 5000
    assert s["by_model"]["claude-opus-5"]["calls"] == 2
    assert s["by_step"]["STEP 5"]["calls"] == 1
    assert s["by_step"]["STEP 6.5"]["calls"] == 2
    sw.reset_usage()


def test_an_unpriced_call_flags_the_bucket_rather_than_inventing_a_cost():
    sw.reset_usage()
    sw.record_usage("anthropic", "claude-sonnet-5", 1000, 100)
    sw.record_usage("gemini", "gemini-not-in-the-table", 5000, 900)
    s = sw.usage_summary()
    assert s["total"]["priced"] is False
    assert s["total"]["input_tokens"] == 6000       # tokens still counted
    assert s["by_model"]["claude-sonnet-5"]["priced"] is True
    sw.reset_usage()


def test_reset_clears_the_ledger():
    sw.reset_usage()
    sw.record_usage("anthropic", "claude-sonnet-5", 10, 10)
    sw.reset_usage()
    assert sw.usage_summary()["total"]["calls"] == 0


def test_usage_files_are_written_and_appended():
    sw.reset_usage()
    sw.log("STEP 1", "research")
    sw.record_usage("anthropic", "claude-sonnet-5", 1000, 500)
    with tempfile.TemporaryDirectory() as d:
        out = sw.Path(d)
        path = sw.write_usage("first-slug", out, "First Article")
        assert path.exists()
        saved = json.loads(path.read_text(encoding="utf-8"))
        assert saved["total"]["calls"] == 1

        sw.record_usage("anthropic", "claude-opus-5", 400, 100)
        sw.write_usage("second-slug", out, "Second Article")

        lines = (out / "usage.jsonl").read_text(encoding="utf-8").strip().split("\n")
        assert len(lines) == 2, lines
        rows = [json.loads(l) for l in lines]
        assert rows[0]["slug"] == "first-slug"
        assert rows[1]["calls"] == 2          # the ledger accumulates across a run
        assert rows[1]["cost_usd"] > rows[0]["cost_usd"]
    sw.reset_usage()


# --- the app ---------------------------------------------------------------

def make_app(tmp, password=""):
    os.environ["APP_PASSWORD"] = password
    os.environ["APP_USERNAME"] = "admin"
    for mod in [m for m in list(sys.modules) if m == "app"]:
        del sys.modules[mod]
    import app as app_module
    app_module.OUTPUT_DIR = sw.Path(tmp)
    app_module.APP_PASSWORD = password
    app_module.app.config["TESTING"] = True
    return app_module


def creds(user="admin", pw="s3cret"):
    raw = base64.b64encode(f"{user}:{pw}".encode()).decode()
    return {"Authorization": f"Basic {raw}"}


def test_no_password_set_leaves_the_app_open():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="")
        assert m.app.test_client().get("/usage").status_code == 200


def test_password_set_blocks_anonymous_requests():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        c = m.app.test_client()
        # Pages send a browser to the login form; APIs answer with a 401.
        for path in ("/", "/usage"):
            r = c.get(path)
            assert r.status_code == 302, (path, r.status_code)
            assert "/login" in r.headers["Location"], path
        for path in ("/api/articles", "/api/usage"):
            r = c.get(path)
            assert r.status_code == 401, (path, r.status_code)
            assert "Basic" in r.headers.get("WWW-Authenticate", ""), path


def test_starting_a_job_requires_the_password():
    # This is the endpoint that spends money; it must never be open.
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        r = m.app.test_client().post("/api/start", json={"topic": "free money"})
        assert r.status_code == 401, r.status_code


def test_correct_password_gets_through():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        r = m.app.test_client().get("/usage", headers=creds())
        assert r.status_code == 200, r.status_code


def test_wrong_password_and_wrong_user_are_both_rejected():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        c = m.app.test_client()
        # On an API route a bad credential is a clean 401 either way.
        assert c.get("/api/usage", headers=creds(pw="guess")).status_code == 401
        assert c.get("/api/usage", headers=creds(user="root")).status_code == 401
        # On a page it falls through to the login form rather than serving it.
        assert c.get("/usage", headers=creds(pw="guess")).status_code == 302


def test_browser_is_sent_to_the_login_page_not_a_popup():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        r = m.app.test_client().get("/usage")
        assert r.status_code == 302, r.status_code
        assert "/login" in r.headers["Location"], r.headers["Location"]
        assert "next=" in r.headers["Location"], r.headers["Location"]


def test_api_callers_still_get_a_401_they_can_handle():
    # A redirect to HTML is useless to a script; /api/* must stay a 401.
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        c = m.app.test_client()
        assert c.get("/api/usage").status_code == 401
        assert c.post("/api/start", json={"topic": "x"}).status_code == 401


def test_login_page_renders_and_is_reachable_without_a_password():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        r = m.app.test_client().get("/login")
        assert r.status_code == 200
        assert b'name="password"' in r.data


def test_signing_in_sets_a_session_and_unlocks_the_app():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        c = m.app.test_client()
        r = c.post("/login", data={"password": "s3cret"})
        assert r.status_code == 302, r.status_code
        # The cookie now carries the session; no basic auth header needed.
        assert c.get("/usage").status_code == 200
        assert c.get("/").status_code == 200


def test_wrong_password_on_the_form_is_rejected():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        c = m.app.test_client()
        r = c.post("/login", data={"password": "nope"})
        assert r.status_code == 401, r.status_code
        assert b"not right" in r.data
        assert c.get("/usage").status_code == 302, "still locked out"


def test_sign_in_returns_you_to_where_you_were_going():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        c = m.app.test_client()
        r = c.post("/login", data={"password": "s3cret", "next": "/usage"})
        assert r.headers["Location"].endswith("/usage"), r.headers["Location"]


def test_next_cannot_bounce_you_to_another_host():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        c = m.app.test_client()
        for hostile in ("//evil.example.com", "https://evil.example.com/x"):
            r = c.post("/login", data={"password": "s3cret", "next": hostile})
            assert r.headers["Location"].endswith("/"), (hostile, r.headers["Location"])
            c.get("/logout")


def test_signing_out_locks_the_app_again():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        c = m.app.test_client()
        c.post("/login", data={"password": "s3cret"})
        assert c.get("/usage").status_code == 200
        c.get("/logout")
        assert c.get("/usage").status_code == 302


def test_basic_auth_still_works_alongside_the_form():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        r = m.app.test_client().get("/api/usage", headers=creds())
        assert r.status_code == 200, r.status_code


def test_repeated_failures_lock_the_attacker_out():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        m._login_failures.clear()
        c = m.app.test_client()
        for _ in range(m._LOGIN_MAX_ATTEMPTS):
            c.post("/login", data={"password": "guess"})
        r = c.post("/login", data={"password": "guess"})
        assert b"Too many attempts" in r.data, r.data[:200]
        # And the lockout holds even against the correct password.
        r = c.post("/login", data={"password": "s3cret"})
        assert b"Too many attempts" in r.data
        m._login_failures.clear()


def test_a_successful_sign_in_clears_the_failure_count():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        m._login_failures.clear()
        c = m.app.test_client()
        for _ in range(3):
            c.post("/login", data={"password": "guess"})
        c.post("/login", data={"password": "s3cret"})
        assert not any(m._login_failures.values()), m._login_failures


def test_login_page_redirects_home_when_no_password_is_configured():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="")
        assert m.app.test_client().get("/login").status_code == 302


def test_the_session_key_changes_with_the_password():
    # Changing the password must invalidate sessions signed under the old one.
    with tempfile.TemporaryDirectory() as d:
        first = make_app(d, password="one").app.secret_key
        second = make_app(d, password="two").app.secret_key
        assert first != second


def test_healthz_stays_open_and_reports_whether_a_password_is_set():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        r = m.app.test_client().get("/healthz")
        assert r.status_code == 200, r.status_code
        assert r.get_json() == {"ok": True, "protected": True}


# --- the dashboard ---------------------------------------------------------

def seed(tmp, rows):
    with open(sw.Path(tmp) / "usage.jsonl", "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r) + "\n")


ROW_A = {"at": "2026-09-01T10:00:00Z", "slug": "a", "title": "Article A",
         "calls": 10, "input_tokens": 50000, "output_tokens": 8000,
         "total_tokens": 58000, "cost_usd": 0.5, "fully_priced": True,
         "by_model": {"claude-sonnet-5": {"calls": 8, "input_tokens": 40000,
                                          "output_tokens": 6000, "cost_usd": 0.2},
                      "claude-opus-5": {"calls": 2, "input_tokens": 10000,
                                        "output_tokens": 2000, "cost_usd": 0.3}}}
ROW_B = {"at": "2026-09-02T11:00:00Z", "slug": "b", "title": "Article B",
         "calls": 5, "input_tokens": 20000, "output_tokens": 3000,
         "total_tokens": 23000, "cost_usd": 0.25, "fully_priced": False,
         "by_model": {"claude-sonnet-5": {"calls": 5, "input_tokens": 20000,
                                          "output_tokens": 3000, "cost_usd": 0.25}}}


def test_rollup_totals_across_runs():
    with tempfile.TemporaryDirectory() as d:
        seed(d, [ROW_A, ROW_B])
        m = make_app(d)
        roll = m.usage_rollup(m.read_usage())
        t = roll["total"]
        assert t["runs"] == 2
        assert t["calls"] == 15
        assert t["total_tokens"] == 81000
        assert abs(t["cost_usd"] - 0.75) < 1e-9
        assert t["fully_priced"] is False        # ROW_B had an unpriced model
        assert roll["by_model"]["claude-sonnet-5"]["calls"] == 13
        assert set(roll["by_day"]) == {"2026-09-01", "2026-09-02"}


def test_newest_run_is_listed_first():
    with tempfile.TemporaryDirectory() as d:
        seed(d, [ROW_A, ROW_B])
        m = make_app(d)
        assert [r["slug"] for r in m.read_usage()] == ["b", "a"]


def test_a_torn_line_does_not_blank_the_dashboard():
    with tempfile.TemporaryDirectory() as d:
        with open(sw.Path(d) / "usage.jsonl", "w", encoding="utf-8") as f:
            f.write(json.dumps(ROW_A) + "\n")
            f.write('{"at": "2026-09-02", "calls": ')      # interrupted write
        m = make_app(d)
        assert len(m.read_usage()) == 1


def test_dashboard_renders_with_no_data():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        r = m.app.test_client().get("/usage")
        assert r.status_code == 200
        assert b"No usage logged yet" in r.data


def test_dashboard_renders_the_numbers():
    with tempfile.TemporaryDirectory() as d:
        seed(d, [ROW_A, ROW_B])
        m = make_app(d)
        body = m.app.test_client().get("/usage").data.decode()
        assert "81,000" in body, "total tokens should be formatted"
        assert "$0.75" in body, "total cost should be shown"
        assert "claude-opus-5" in body
        assert "Article A" in body and "Article B" in body
        assert "Partial pricing" in body, "the unpriced warning should appear"


def test_usage_api_returns_rollup_and_runs():
    with tempfile.TemporaryDirectory() as d:
        seed(d, [ROW_A, ROW_B])
        m = make_app(d)
        data = m.app.test_client().get("/api/usage").get_json()
        assert data["total"]["runs"] == 2
        assert len(data["runs"]) == 2
        assert "by_model" in data and "by_day" in data


# --- evaluating a draft from the browser -----------------------------------

REVIEW_RECORD = {
    "agents": {"writer": {"model": "claude-sonnet-5", "provider": "anthropic"},
               "auditor": {"model": "gpt-5.5", "provider": "openai"},
               "judge": {"model": "claude-opus-5", "provider": "anthropic"}},
    "outcome": "passed on round 2",
    "rounds": [{
        "round": 1, "verdict": "revise", "scores": {"factual_support": 6},
        "issues": [
            {"id": "i1", "category": "factual", "severity": "high",
             "quote": "cuts p99 latency by 40%", "problem": "No citation.",
             "fix": "Attribute it or cut it."},
            {"id": "i2", "category": "ai_tell", "severity": "low",
             "quote": "a pivotal moment", "problem": "Inflation.", "fix": "Cut it."},
        ],
        "responses": [{"id": "i1", "stance": "dispute", "reason": "cited below"}],
        "rulings": [{"id": "i1", "ruling": "uphold", "reasoning": "different claim"}],
        "applied": ["i1", "i2"],
    }],
}


def with_review(tmp, slug="demo"):
    m = make_app(tmp)
    (sw.Path(tmp) / f"{slug}_review.json").write_text(
        json.dumps(REVIEW_RECORD), encoding="utf-8")
    return m


def test_audit_needs_something_to_audit():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        r = m.app.test_client().post("/api/audit/start", data={})
        assert r.status_code == 400, r.status_code
        assert "Paste an article" in r.get_json()["error"]


def test_audit_rejects_a_draft_too_short_to_be_worth_the_calls():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        r = m.app.test_client().post("/api/audit/start", data={"text": "far too short"})
        assert r.status_code == 400
        assert "too short" in r.get_json()["error"]


def test_audit_rejects_file_types_it_cannot_read():
    import io
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        r = m.app.test_client().post(
            "/api/audit/start",
            data={"file": (io.BytesIO(b"binary"), "payload.exe")},
            content_type="multipart/form-data")
        assert r.status_code == 400
        assert ".exe is not supported" in r.get_json()["error"], r.get_json()


def test_audit_rejects_a_file_that_is_not_utf8():
    import io
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        r = m.app.test_client().post(
            "/api/audit/start",
            data={"file": (io.BytesIO(b"\xff\xfe\x00binary"), "draft.md")},
            content_type="multipart/form-data")
        assert r.status_code == 400
        assert "not UTF-8" in r.get_json()["error"], r.get_json()


def test_audit_is_behind_the_password_like_everything_else():
    # It spends money, so it must never be open.
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d, password="s3cret")
        r = m.app.test_client().post("/api/audit/start", data={"text": "x " * 100})
        assert r.status_code == 401, r.status_code


def test_review_page_shows_all_three_voices():
    with tempfile.TemporaryDirectory() as d:
        m = with_review(d)
        body = m.app.test_client().get("/review/demo").data.decode()
        assert "cuts p99 latency by 40%" in body        # the quoted text
        assert "gpt-5.5" in body                        # the roster
        assert "disputed" in body                       # the writer
        assert "uphold" in body                         # the judge
        assert "did not respond" in body                # i2, never answered
        assert "passed on round 2" in body              # the outcome


def test_review_counts_are_computed_not_guessed():
    with tempfile.TemporaryDirectory() as d:
        m = with_review(d)
        body = m.app.test_client().get("/review/demo").data.decode()
        assert "<b>2</b> raised" in body, body[body.find("counts"):][:300]
        assert "<b>1</b> disputed" in body
        assert "<b>2</b> applied" in body


def test_missing_review_is_a_404_not_a_crash():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        r = m.app.test_client().get("/review/nothing-here")
        assert r.status_code == 404
        assert b"No review found" in r.data


def test_review_slug_cannot_escape_the_output_directory():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        r = m.app.test_client().get("/review/..%2f..%2fetc%2fpasswd")
        assert r.status_code == 404, r.status_code


def test_reviews_are_listed_with_their_headline_numbers():
    with tempfile.TemporaryDirectory() as d:
        m = with_review(d)
        rows = m.app.test_client().get("/api/reviews").get_json()
        assert len(rows) == 1, rows
        assert rows[0]["slug"] == "demo"
        assert rows[0]["raised"] == 2 and rows[0]["applied"] == 2
        assert rows[0]["agents"]["auditor"] == "gpt-5.5"


def test_the_page_offers_both_modes():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        body = m.app.test_client().get("/").data.decode()
        assert "Evaluate a draft" in body
        assert "Generate New Article" in body


def test_docx_headings_survive_the_conversion():
    try:
        from docx import Document
    except ImportError:
        return                      # python-docx absent; nothing to check
    import io
    with tempfile.TemporaryDirectory() as d:
        make_app(d)
        doc = Document()
        doc.add_heading("The Title", level=1)
        doc.add_paragraph("Some body text.")
        doc.add_heading("A Section", level=2)
        doc.add_paragraph("More body text.")
        buf = io.BytesIO()
        doc.save(buf)
        import app as m
        md = m._docx_to_markdown(buf.getvalue())
        assert "# The Title" in md, md
        assert "## A Section" in md, md
        assert "Some body text." in md


# --- the job stream ---------------------------------------------------------

def drain(m, job_id, limit=40):
    """Collect SSE frames from the stream endpoint."""
    r = m.app.test_client().get(f"/api/stream/{job_id}")
    out = []
    for chunk in r.response:
        out.append(chunk.decode() if isinstance(chunk, bytes) else chunk)
        if len(out) >= limit:
            break
    return "".join(out)


def test_a_failure_is_sent_as_failed_not_error():
    # "error" is reserved by EventSource for connection trouble; sending
    # pipeline failures under it made a dropped connection look like a crash.
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        q = m.queue.Queue()
        with m._jobs_lock:
            m._jobs["j1"] = q
            m._job_started["j1"] = m.time.time()
        q.put(("log", "starting"))
        q.put(("error", "the model hit its cap"))
        body = drain(m, "j1")
    assert "event: failed" in body, body
    assert "event: error" not in body, body
    assert "the model hit its cap" in body


def test_a_finished_job_is_forgotten():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        q = m.queue.Queue()
        with m._jobs_lock:
            m._jobs["j2"] = q
            m._job_started["j2"] = m.time.time()
        q.put(("done", '{"review_slug": "x"}'))
        body = drain(m, "j2")
        assert "event: done" in body, body
        assert "j2" not in m._jobs, "a completed job should not be retained"


def test_an_unknown_job_is_a_404():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        assert m.app.test_client().get("/api/stream/nope").status_code == 404


def test_abandoned_jobs_are_swept_but_recent_ones_are_kept():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        with m._jobs_lock:
            m._jobs["old"] = m.queue.Queue()
            m._job_started["old"] = m.time.time() - m.JOB_RETENTION_SECS - 60
            m._jobs["new"] = m.queue.Queue()
            m._job_started["new"] = m.time.time()
        m._reap_jobs()
        assert "old" not in m._jobs, "an abandoned job should be swept"
        assert "new" in m._jobs, "a live job must survive the sweep"


def test_the_client_listens_for_failed_and_handles_a_drop_separately():
    with tempfile.TemporaryDirectory() as d:
        m = make_app(d)
        body = m.app.test_client().get("/").data.decode()
        assert "addEventListener('failed'" in body, "the client must listen for failed"
        assert "Reconnecting" in body, "a drop should say it is reconnecting"
        assert "'Pipeline error'" not in body, \
            "the hardcoded fallback string should be gone"


CASES = [v for k, v in sorted(globals().items()) if k.startswith("test_")]


def main():
    saved = {k: os.environ.get(k) for k in ("APP_PASSWORD", "APP_USERNAME")}
    failures = []
    for case in CASES:
        try:
            case()
        except Exception as e:
            failures.append((case.__name__, "%s: %s" % (type(e).__name__, e)))
    for k, v in saved.items():
        if v is None:
            os.environ.pop(k, None)
        else:
            os.environ[k] = v
    print("\n" + "=" * 60)
    for name, err in failures:
        print("FAIL %s\n     %s" % (name, err))
    print("%d/%d passed" % (len(CASES) - len(failures), len(CASES)))
    print("=" * 60)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
